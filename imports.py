"""Strict, versioned synthetic adapters; preview and commit use the same validation."""
import csv
import hashlib
import io
import math
import zipfile
from datetime import date
from decimal import Decimal, InvalidOperation
from openpyxl import Workbook, load_workbook
from domain import COUNTRIES, OPERATORS, OPERATOR_CODES, PROFILE_FEES, rebuild

FIELDS = {
 'inventory': ['iccid','profile_id','department','activated'],
 'usage': ['iccid','date','country','data_mb','minutes','sms'],
 'invoice': ['iccid','month','billed_eur']}
HEADERS = {
 'Movistar': {'inventory':['sim_id','plan_code','business_unit','activation_date'],
   'usage':['sim_id','usage_date','country_code','data_mb','voice_minutes','sms_count'],
   'invoice':['sim_id','billing_month','total_eur']},
 'Vodafone': {'inventory':['ICCID','Tariff','Department','Activated'],
   'usage':['ICCID','Day','Country','Data_KB','Voice_Seconds','SMS'],
   'invoice':['ICCID','Period','Amount_Cents']},
 'Orange': {'inventory':['line_identifier','offer','cost_centre','start_date'],
   'usage':['line_identifier','event_date','destination','volume_gb','call_minutes','messages'],
   'invoice':['line_identifier','invoice_period','amount_eur']}}
MAX_ROWS=5000

def parse_file(raw, filename, operator, kind, db):
    if operator not in OPERATORS or kind not in FIELDS:
        raise ValueError('Choose a supported operator and document type.')
    expected = HEADERS[operator][kind]
    if filename.lower().endswith('.csv'):
        try:
            reader=csv.reader(io.StringIO(raw.decode('utf-8-sig')),delimiter=';' if operator=='Orange' else ',')
            headers=next(reader,[]); rows=[]
            for r in reader:
                rows.append(r)
                if len(rows)>MAX_ROWS: raise ValueError('Maximum 5.000 rows per demo import.')
        except UnicodeDecodeError:
            raise ValueError('CSV must use UTF-8 encoding.')
    elif filename.lower().endswith('.xlsx'):
        with zipfile.ZipFile(io.BytesIO(raw)) as archive:
            if len(archive.infolist())>200:
                raise ValueError('Workbook contains too many components.')
            if sum(entry.file_size for entry in archive.infolist()) > 25*1024*1024:
                raise ValueError('Expanded workbook exceeds the 25 MB demo limit.')
            if any('vbaproject' in entry.filename.lower() or 'externallinks/' in entry.filename.lower()
                   or '/embeddings/' in entry.filename.lower() for entry in archive.infolist()):
                raise ValueError('Workbooks with macros, external links or embedded files are not accepted.')
        book=load_workbook(io.BytesIO(raw),read_only=True,data_only=False,keep_links=False)
        try:
            sheet=book.active
            iterator=sheet.iter_rows(values_only=True)
            headers=list(next(iterator,[])); rows=[]
            for r in iterator:
                rows.append(list(r))
                if len(rows)>MAX_ROWS: raise ValueError('Maximum 5.000 rows per demo import.')
        finally:
            book.close()
    else:
        raise ValueError('Use a .csv or .xlsx file.')
    if headers!=expected:
        raise ValueError('The column names do not match the selected format. Expected: '+', '.join(expected))
    if not rows: raise ValueError('The file contains no data rows.')
    normalized=[]; errors=[]; seen=set()
    sims={r['iccid']:r['operator'] for r in db.execute('SELECT iccid,operator FROM sims')}
    profiles={r['id']:r['operator'] for r in db.execute('SELECT id,operator FROM profiles')}
    activated={r['iccid']:r['activated'] for r in db.execute('SELECT iccid,activated FROM sims')}
    for idx,r in enumerate(rows,2):
        try:
            if len(r)!=len(expected): raise ValueError('Incorrect number of columns.')
            if any(isinstance(v,str) and v.startswith('=') for v in r): raise ValueError('Spreadsheet formulas are not accepted.')
            if not isinstance(r[0],str): raise ValueError('ICCID must be text to preserve all 19 digits.')
            item=dict(zip(FIELDS[kind], [str(v).strip() if v is not None else '' for v in r]))
            iccid=item['iccid']
            if len(iccid)!=19 or not iccid.isascii() or not iccid.isdigit(): raise ValueError('ICCID must contain exactly 19 digits.')
            if kind=='inventory':
                if iccid in sims: raise ValueError('This SIM already exists. Inventory files can only add new SIMs.')
                if profiles.get(item['profile_id'])!=operator: raise ValueError('Profile does not belong to selected operator.')
                if not 1<=len(item['department'])<=60: raise ValueError('Department must contain 1–60 characters.')
                stamp=item['activated']
            else:
                if sims.get(iccid)!=operator: raise ValueError('The SIM must exist under the selected operator. Check the operator or import its inventory first.')
                stamp=item['date'] if kind=='usage' else item['month']+'-01'
            if date.fromisoformat(stamp).isoformat()!=stamp: raise ValueError('Use dates in YYYY-MM-DD format.')
            if not '2026-01-01'<=stamp<='2026-08-31': raise ValueError('Demo dates must fall between January and August 2026.')
            if kind=='usage':
                if stamp<activated[iccid]: raise ValueError('The usage date is before the SIM activation date.')
                item['country']=item['country'].upper()
                if item['country'] not in COUNTRIES: raise ValueError('Unsupported country code for this demo.')
                for field in ('data_mb','minutes','sms'):
                    value=float(item[field].replace(',','.') if operator=='Orange' else item[field])
                    if not math.isfinite(value) or value<0: raise ValueError('Usage must be finite and non-negative.')
                    item[field]=value
                if operator=='Vodafone': item['data_mb']/=1024; item['minutes']/=60
                if operator=='Orange': item['data_mb']*=1024
                if item['sms']!=int(item['sms']): raise ValueError('SMS must be a whole number.')
                if item['data_mb']>1_000_000 or item['minutes']>1440 or item['sms']>100000:
                    raise ValueError('Daily usage exceeds demo safety limits.')
                item['sms']=int(item['sms']); item['data_mb']=round(item['data_mb'],6)
                key=(iccid,stamp,item['country'])
                exists=db.execute('SELECT 1 FROM usage WHERE iccid=? AND date=? AND country=?',key).fetchone()
            elif kind=='invoice':
                if stamp[:7]<activated[iccid][:7]: raise ValueError('The invoice month is before the SIM activation month.')
                amount=Decimal(item['billed_eur'].replace(',','.') if operator=='Orange' else item['billed_eur'])
                cents=amount if operator=='Vodafone' else amount*100
                if not cents.is_finite() or not 0<=cents<=100_000_000 or cents!=cents.to_integral_value():
                    raise ValueError('Amount must be non-negative with cent precision.')
                item['billed_cents']=int(cents); del item['billed_eur']
                key=(iccid,item['month'])
                exists=db.execute('SELECT 1 FROM invoices WHERE iccid=? AND month=?',key).fetchone()
            else:
                key=(iccid,); exists=False
            if key in seen or exists: raise ValueError('This record already exists in the file or database. Existing data will not be replaced.')
            seen.add(key); normalized.append(item)
        except (ValueError,InvalidOperation,OverflowError) as error:
            errors.append({'row':idx,'message':str(error) or 'Invalid value.'})
    digest=hashlib.sha256(operator.encode()+kind.encode()+raw).hexdigest()
    if db.execute('SELECT 1 FROM imports WHERE digest=?',(digest,)).fetchone():
        errors.insert(0,{'row':0,'message':'This exact file was already imported.'})
    return {'rows':normalized,'errors':errors,'total':len(rows),'valid':len(normalized),'digest':digest}

def commit(db,result,filename,operator,kind):
    if result['errors']: raise ValueError('Fix all the reported errors before importing.')
    source='import:'+result['digest'][:12]
    rows=result['rows']; months=set()
    for r in rows:
        if kind=='inventory':
            db.execute('INSERT INTO sims VALUES (?,?,?,?,?,?)',(r['iccid'],operator,r['profile_id'],r['department'],r['activated'],'Active'))
            months.update(f'2026-{m:02d}' for m in range(int(r['activated'][5:7]),9))
        elif kind=='usage':
            db.execute('INSERT INTO usage VALUES (?,?,?,?,?,?,?)',(r['iccid'],r['date'],r['country'],r['data_mb'],r['minutes'],r['sms'],source))
            months.add(r['date'][:7])
        else:
            db.execute('INSERT INTO invoices VALUES (?,?,?,?)',(r['iccid'],r['month'],r['billed_cents'],source))
    if months: rebuild(db,sorted(months))
    db.execute('INSERT INTO imports(filename,operator,kind,rows,digest) VALUES (?,?,?,?,?)',
               (filename,operator,kind,len(rows),result['digest']))

def sample(operator,kind,extension):
    op=OPERATORS.index(operator)
    rows=[]
    for n in range(3):
        iccid=f'893499{op*100+n:013d}'
        if kind=='inventory': r=[iccid,OPERATOR_CODES[operator]+'-500','Demo import team','2026-08-01']
        elif kind=='usage':
            r=[iccid,'2026-08-15','AD',50,5,2]
            if operator=='Vodafone': r[3]*=1024; r[4]*=60
            if operator=='Orange': r[3]/=1024
        else:
            # 50MB / 5min / 2SMS in Andorra plus the full monthly fee.
            cents=PROFILE_FEES[operator][500]+round((50*2.8+5*65+2*18)*[1,.92,1.08][op])
            r=[iccid,'2026-08',cents if operator=='Vodafone' else f'{cents/100:.2f}']
        rows.append(r)
    headers=HEADERS[operator][kind]
    if extension=='xlsx':
        book=Workbook(); sheet=book.active; sheet.title=kind.title(); sheet.append(headers)
        for row in rows: sheet.append(row)
        for cell in sheet['A']: cell.number_format='@'
        out=io.BytesIO(); book.save(out); return out.getvalue()
    out=io.StringIO(newline=''); writer=csv.writer(out,delimiter=';' if operator=='Orange' else ',')
    writer.writerow(headers); writer.writerows(rows); return out.getvalue().encode('utf-8-sig')
